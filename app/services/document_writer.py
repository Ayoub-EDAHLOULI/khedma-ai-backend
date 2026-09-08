"""Generates downloadable .docx files for a tailored resume and cover letter.

Two strategies for the resume:
- Layout-preserving edit: when the profile has a stored original .docx, ask the LLM
  to rewrite each non-empty paragraph in place (same count/order), then substitute
  that text into the original document's paragraphs — this keeps the original fonts,
  bullets, and headers because we never touch paragraph/run formatting, only text.
  Only reliable for simple, single-column resumes (section 12 in the docs) — if the
  LLM's paragraph count doesn't match the original, we abandon the edit and fall back.
- Fresh generation: used whenever there's no stored original (or the edit above was
  abandoned) — a clean, simple ATS-style single-column document built from scratch.

The cover letter has no "original" to preserve, so it's always freshly generated.
"""
import io
import json

from docx import Document
from openai import OpenAI

from app.config import settings
from app.models import Job, Profile

MODEL = "gpt-4o-mini"

_client = OpenAI(api_key=settings.openai_api_key)


def _non_empty_paragraphs(document: Document) -> list[int]:
    return [i for i, p in enumerate(document.paragraphs) if p.text.strip()]


def _rewrite_paragraphs_aligned(
    original_texts: list[str], tailored_cv: str, job: Job
) -> list[str] | None:
    """Ask the LLM to rewrite each paragraph in place, foregrounding what's relevant
    to this job. Returns None if the model didn't return exactly one line per input
    paragraph (paragraph-count mismatch — layout-preserving edit isn't safe)."""
    system_prompt = (
        "You are editing a candidate's resume paragraph by paragraph. You will "
        "receive a JSON array of the resume's existing paragraphs (in order) and "
        "a target job. Rewrite EACH paragraph to better foreground relevant "
        "experience/skills for this job — reorder emphasis within a paragraph, "
        "tighten wording — but do not invent new experience, skills, dates, or "
        "employers not already present in the candidate's own text. If a "
        "paragraph is a section header (e.g. 'Experience', 'Skills') or contact "
        "info, return it unchanged. Respond ONLY with a JSON object: "
        '{"paragraphs": ["...", ...]} with EXACTLY one string per input paragraph, '
        "in the same order — never merge, split, or drop paragraphs."
    )
    user_prompt = json.dumps(
        {
            "paragraphs": original_texts,
            "job_title": job.title,
            "company": job.company,
            "job_description": (job.description or "")[:2000],
            "already_tailored_summary": tailored_cv,
        }
    )

    response = _client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
        temperature=0.2,
    )

    rewritten = json.loads(response.choices[0].message.content).get("paragraphs")
    if not isinstance(rewritten, list) or len(rewritten) != len(original_texts):
        return None
    return rewritten


def _set_paragraph_text(paragraph, text: str) -> None:
    """Replace a paragraph's visible text while keeping its first run's formatting
    (font, bold, size, etc.) and the paragraph's own style (bullets, headings)."""
    if not paragraph.runs:
        paragraph.add_run(text)
        return

    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def tailor_resume_docx(profile: Profile, tailored_cv: str, job: Job) -> bytes:
    if profile.resume_docx:
        edited = _try_layout_preserving_edit(profile, tailored_cv, job)
        if edited is not None:
            return edited

    return _generate_fresh_resume_docx(profile, tailored_cv)


def _try_layout_preserving_edit(
    profile: Profile, tailored_cv: str, job: Job
) -> bytes | None:
    document = Document(io.BytesIO(profile.resume_docx))
    indices = _non_empty_paragraphs(document)
    if not indices:
        return None

    original_texts = [document.paragraphs[i].text for i in indices]
    rewritten = _rewrite_paragraphs_aligned(original_texts, tailored_cv, job)
    if rewritten is None:
        return None

    for i, new_text in zip(indices, rewritten):
        _set_paragraph_text(document.paragraphs[i], new_text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _generate_fresh_resume_docx(profile: Profile, tailored_cv: str) -> bytes:
    document = Document()

    document.add_heading(profile.full_name or "Resume", level=1)
    if profile.skills:
        document.add_heading("Skills", level=2)
        document.add_paragraph(", ".join(profile.skills))

    document.add_heading("Experience", level=2)
    for line in tailored_cv.split("\n"):
        line = line.strip()
        if not line:
            continue
        document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def generate_cover_letter_docx(cover_letter: str, profile: Profile, job: Job) -> bytes:
    document = Document()

    document.add_heading(profile.full_name or "Cover Letter", level=1)
    document.add_paragraph(f"Re: {job.title}" + (f" — {job.company}" if job.company else ""))
    document.add_paragraph("")

    for paragraph in cover_letter.split("\n\n"):
        paragraph = paragraph.strip()
        if paragraph:
            document.add_paragraph(paragraph)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
